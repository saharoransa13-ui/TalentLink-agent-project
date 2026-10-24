
from io import BytesIO
from pathlib import Path
from pypdf import PdfReader
from docx import Document


MAX_FILE_SIZE = 5 * 1024 * 1024  # 5 MB
MAX_PAGES = 10

def extract_cv_text(filename: str, content: bytes) -> str:
    extension = Path(filename).suffix.lower()

    if extension not in {".pdf", ".docx"}:
        raise ValueError("Only PDF and DOCX files are supported.")

    if len(content) > MAX_FILE_SIZE:
        raise ValueError("The file must not exceed 5 MB.")

    if extension == ".pdf":
        reader = PdfReader(BytesIO(content))

        if len(reader.pages) > MAX_PAGES:
            raise ValueError("The PDF must not exceed 10 pages.")

        if reader.is_encrypted:
            raise ValueError("Encrypted PDF files are not supported.")

        cv_text = "\n".join(
            page.extract_text() or ""
            for page in reader.pages
        )

    else:
        document = Document(BytesIO(content))

        paragraphs = [
            paragraph.text
            for paragraph in document.paragraphs
        ]

        for table in document.tables:
            for row in table.rows:
                paragraphs.append(
                    " | ".join(cell.text for cell in row.cells)
                )

        cv_text = "\n".join(paragraphs)

    if not cv_text.strip():
        raise ValueError(
            "No readable text found. The file may need OCR."
        )

    return cv_text
